from pdf2docx import Converter
from docx import Document
from pathlib import Path

def pdf_to_docx(input_path: str, output_path: str):
    cv = Converter(input_path)
    cv.convert(output_path)
    cv.close()

def docx_to_pdf(input_path: str, output_path: str):
    from fpdf import FPDF
    doc = Document(input_path)
    pdf = FPDF()
    pdf.add_page()
    for paragraph in doc.paragraphs:
        pdf.set_font("Arial", size=12)
        pdf.cell(200, 10, txt=paragraph.text, ln=True)
    pdf.output(output_path)